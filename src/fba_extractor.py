class FBAExtractor:
    def __init__(self, tables, docx_extractor = None):
        self.tables = tables
        self.docx_extractor = docx_extractor

    def normalize_label(self, text):
        text = text.strip().lower()
        text = text.rstrip(":")
        text = " ".join(text.split())

        return text

    def find_table(self, possible_labels):
        normalized_labels = {
            self.normalize_label(label)
            for label in possible_labels
        }

        for table in self.tables:
            for row in table:
                for cell in row:
                    normalized_cell = self.normalize_label(cell)

                    for label in normalized_labels:
                        if (
                            normalized_cell == label
                            or normalized_cell.startswith(label + ":")
                            or normalized_cell.startswith(label + " ")
                        ):
                            return table

        return None

    def find_value_after_labels(self, table, possible_labels, all_labels):
        normalized_possible_labels = {
            self.normalize_label(label)
            for label in possible_labels
        }

        normalized_all_labels = {
            self.normalize_label(label)
            for label in all_labels
        }

        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                normalized_cell = self.normalize_label(cell)

                matched_label = any(
                    normalized_cell == label
                    or normalized_cell.startswith(label + ":")
                    for label in normalized_possible_labels
                )

                if matched_label:

                    # Case 1: value is in the next column
                    if column_index + 1 < len(row):
                        value = row[column_index + 1].strip()

                        if (
                            value
                            and self.normalize_label(value)
                            not in normalized_all_labels
                        ):
                            return value

                    # Case 2: value is in the same column on the next row
                    if row_index + 1 < len(table):
                        next_row = table[row_index + 1]

                        if column_index < len(next_row):
                            value = next_row[column_index].strip()

                            if (
                                value
                                and self.normalize_label(value)
                                not in normalized_all_labels
                            ):
                                return value

        return None

    def extract_identification(self):
        table = self.find_table([
        "Member Name",
        "Client Name",
        "Patient Name",])

        if table is None:
            return {}

        field_labels = {
            "member_name": [
                "Member Name",
                "Client Name",
                "Patient Name",
            ],
            "member_dob": [
                "Member DOB",
                "Client DOB",
                "Patient DOB",
                "Date of Birth",
                "DOB",
            ],
            "cin": [
                "CIN #",
                "CIN",
                "Client Identification Number",
            ],
            "diagnosis": [
                "Diagnoses/with ICD Code",
                "Diagnosis",
                "Diagnosis with ICD Code",
                "ICD Diagnosis",
            ],
            "guardian_name": [
                "Guardian Name",
                "Parent/Guardian Name",
                "Caregiver Name",
            ],
            "phone": [
                "Phone",
                "Phone Number",
                "Contact Number",
            ],
            "primary_care_provider": [
                "Primary Care Provider",
                "PCP",
            ],
            "known_allergies": [
                "Known Allergies",
                "Allergies",
            ],
            "current_medications_dosage": [
                "Current Medications/Dosage",
                "Current Medications",
                "Medications",
            ],
            "dietary_restrictions": [
                "Dietary Restrictions",
                "Diet Restrictions",
            ],
            "lmhp_name_credential": [
                "Full name & Credential",
                "LMHP Name",
                "LMHP Name & Credential",
            ],
            "lmhp_contact_number": [
                "Contact Number",
                "LMHP Contact Number",
                "LMHP Phone",
            ],
        }

        # Flatten all aliases so the extractor can recognize when
        # the neighboring cell is another label instead of a value.
        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        identification = {}

        for field_name, aliases in field_labels.items():
            identification[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return identification

    
    def extract_service_history(self):
        table = self.find_table([
            "Service Initiation Date",
            "Service Start Date",
            "Date ABA first began",
            "Prior Applied Behavioral Health Agencies",
        ])

        if table is None:
            return {}

        field_labels = {
            "service_initiation_date": [
                "Service Initiation Date",
                "Service Start Date",
                "Start Date",
            ],
            "aba_first_began": [
                "Date ABA first began",
                "ABA Start Date",
                "Date ABA Began",
            ],
            "prior_aba_agencies": [
                "Prior Applied Behavioral Health Agencies",
                "Prior ABA Agencies",
                "Previous ABA Agencies",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        service_history = {}

        # Extract known fields
        for field_name, aliases in field_labels.items():
            service_history[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        # Track values already extracted
        known_values = {
            str(value).strip()
            for value in service_history.values()
            if value is not None
        }

        # Normalize known labels for comparison
        normalized_labels = {
            self.normalize_label(label)
            for label in all_labels
        }

        # Preserve values that could not be classified
        unclassified_values = []

        for row in table:
            for cell in row:
                value = str(cell).strip()

                if not value:
                    continue

                if value in known_values:
                    continue

                if self.normalize_label(value) in normalized_labels:
                    continue

                if value not in unclassified_values:
                    unclassified_values.append(value)

        service_history["unclassified_values"] = unclassified_values

        return service_history


    def extract_administrative_contact(self):
        table = self.find_table([
            "Administrative Contact for Current Authorization Request",
            "Full Name and Title",
            "Name and Title",
        ])

        if table is None:
            return {}

        field_labels = {
            "full_name_title": [
                "Full Name and Title",
                "Name and Title",
                "Administrative Contact",
            ],
            "phone_number": [
                "Phone Number",
                "Phone",
                "Contact Number",
            ],
            "fax_number": [
                "Fax Number",
                "Fax",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        administrative_contact = {}

        for field_name, aliases in field_labels.items():
            administrative_contact[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return administrative_contact

    def extract_chief_complaint(self):
        table = self.find_table([
            "Chief Complaint/Reason for Seeking Applied Behavior Analysis (ABA) Treatment",
            "Chief Complaint",
            "Reason for Seeking ABA Treatment",
            "Reason for Referral",
        ])

        if table is None:
            return {}

        field_labels = {
            "chief_complaint": [
                "Chief Complaint/Reason for Seeking Applied Behavior Analysis (ABA) Treatment",
                "Chief Complaint",
                "Reason for Seeking ABA Treatment",
                "Reason for Referral",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        chief_complaint = {}

        for field_name, aliases in field_labels.items():
            chief_complaint[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return chief_complaint

    def extract_records_reviewed(self):
        table = self.find_table([
            "Record Type",
            "Author of Record",
            "Date of Record",
            "Records Reviewed",])

        if table is None:
            return []
        

        records = []

        for row in table:
            # Skip rows that don't contain enough columns
            if len(row) < 3:
                continue

            record_type = row[0].strip()
            author = row[1].strip()
            date = row[2].strip()

            # Skip merged title rows where every column contains the same text
            non_empty_values = [
                cell.strip()
                for cell in row
                if cell.strip()
            ]

            if (
                non_empty_values
                and len(set(non_empty_values)) == 1
                and "records reviewed" in self.normalize_label(non_empty_values[0])
            ):
                continue

            # Skip the column header row
            if self.normalize_label(record_type) == "record type":
                continue

            # Skip completely empty rows
            if not record_type and not author and not date:
                continue

            records.append({
                "record_type": record_type,
                "author": author,
                "date": date,
            })

        return records

    def extract_interviews_conducted(self):
        table = self.find_table([
            "Initial Interview/Observation",
            "Second Interview/Observation",
            "Interview/Observation",
        ])

        if table is None:
            return []

        interviews = []

        for row_index, row in enumerate(table):
            if not row:
                continue

            label = row[0].strip()

            # Find any interview/observation label
            if "interview/observation" not in label.lower():
                continue

            # Remove instructions after the colon
            observation_type = label.split(":")[0].strip()

            # Get the narrative from the following row
            if row_index + 1 < len(table):
                next_row = table[row_index + 1]

                description = next(
                    (
                        cell.strip()
                        for cell in next_row
                        if cell.strip()
                    ),
                    None
                )

                if description:
                    interviews.append({
                        "observation_type": observation_type,
                        "description": description,
                    })

        return interviews

    def extract_individual_description(self):
        table = self.find_table([
            "Individual Description/Living Arrangements",
            "Individual Description",
            "Living Arrangements",
        ])

        if table is None:
            return {}

        field_labels = {
            "individual_description_living_arrangements": [
                "Individual Description/Living Arrangements",
                "Individual Description",
                "Living Arrangements",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        individual_description = {}

        for field_name, aliases in field_labels.items():
            individual_description[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return individual_description

    def extract_significant_medical_history(self):
        table = self.find_table([
            "Significant Medical History",
            "Medical History",
        ])

        if table is None:
            return {}

        field_labels = {
            "significant_medical_history": [
                "Significant Medical History",
                "Medical History",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        medical_history = {}

        for field_name, aliases in field_labels.items():
            medical_history[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return medical_history

    def extract_functional_communication_skills(self):
        table = self.find_table([
            "Functional Communication Skills",
            "Communication Skills",
        ])

        if table is None:
            return {}

        field_labels = {
            "functional_communication_skills": [
                "Functional Communication Skills",
                "Communication Skills",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        communication_skills = {}

        for field_name, aliases in field_labels.items():
            communication_skills[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return communication_skills

    def extract_self_care_adls(self):
        table = self.find_table([
            "Self-Care and Activities of Daily Living Skills",
            "Self-Care/ADLs",
            "Self-Care",
            "Activities of Daily Living",
            "ADLs",
        ])

        if table is None:
            return {}

        field_labels = {
            "self_care_adls": [
                "Self-Care and Activities of Daily Living Skills",
                "Self-Care/ADLs",
                "Self-Care",
                "Activities of Daily Living",
                "ADLs",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        self_care = {}

        for field_name, aliases in field_labels.items():
            self_care[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return self_care

    def extract_social_play_skills(self):
        table = self.find_table([
            "Social and Play Skills",
            "Social/Play Skills",
            "Social Skills",
            "Play Skills",
        ])

        if table is None:
            return {}

        field_labels = {
            "social_play_skills": [
                "Social and Play Skills",
                "Social/Play Skills",
                "Social Skills",
                "Play Skills",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        social_play_skills = {}

        for field_name, aliases in field_labels.items():
            social_play_skills[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return social_play_skills

    def extract_mobility_functioning_restrictions(self):
        table = self.find_table([
            "Mobility Functioning and Restrictions",
            "Mobility Functioning",
            "Mobility Restrictions",
            "Mobility",
        ])

        if table is None:
            return {}

        field_labels = {
            "mobility_functioning_restrictions": [
                "Mobility Functioning and Restrictions",
                "Mobility Functioning",
                "Mobility Restrictions",
                "Mobility",
            ],
        }

        all_labels = []

        for aliases in field_labels.values():
            all_labels.extend(aliases)

        mobility = {}

        for field_name, aliases in field_labels.items():
            mobility[field_name] = self.find_value_after_labels(
                table,
                aliases,
                all_labels
            )

        return mobility

    def extract_daily_schedule(self):
        table = self.find_table([
            "Daily schedule of all activities",
            "Daily Schedule",
            "Daily Routine",
        ])

        if table is None:
            return {}

        schedule = {}

        days = {
            "monday",
            "tuesday",
            "wednesday",
            "thursday",
            "friday",
            "saturday",
            "sunday",
        }

        for row_index, row in enumerate(table):
            normalized_row = [
                self.normalize_label(cell)
                for cell in row
            ]

            # Find the row containing the days of the week
            if any(cell in days for cell in normalized_row):

                # Schedule information should be in the next row
                if row_index + 1 >= len(table):
                    break

                activity_row = table[row_index + 1]

                for column_index, day in enumerate(row):
                    day_name = self.normalize_label(day)

                    if day_name in days:
                        value = ""

                        if column_index < len(activity_row):
                            value = activity_row[column_index].strip()

                        schedule[day_name] = value

                break

        return schedule
    
    def extract_daily_school_schedule(self):
        table = self.find_table([
            "Daily School Schedule",
            "School Schedule",
        ])

        if table is None:
            return {}

        schedule = {}

        days = {
            "monday",
            "tuesday",
            "wednesday",
            "thursday",
            "friday",
            "saturday",
            "sunday",
        }

        for row_index, row in enumerate(table):
            normalized_row = [
                self.normalize_label(cell)
                for cell in row
            ]

            # Find the row containing the days of the week
            if any(cell in days for cell in normalized_row):

                if row_index + 1 >= len(table):
                    break

                activity_row = table[row_index + 1]

                for column_index, day in enumerate(row):
                    day_name = self.normalize_label(day)

                    if day_name in days:
                        value = ""

                        if column_index < len(activity_row):
                            value = activity_row[column_index].strip()

                        schedule[day_name] = value

                break

        return schedule

    def extract_school_iep_information(self):
        table = self.find_table([
            "Are ABA services being requested for authorization from CalOptima Health at the school setting?",
            "Does the member have a current Individualized Educational Plan",
        ])

        if table is None:
            return {}

        field_labels = {
            "aba_services_requested_at_school": [
                "Are ABA services being requested for authorization from CalOptima Health at the school setting?",
            ],
            "has_current_iep": [
                "Does the member have a current Individualized Educational Plan (IEP/equivalent)?",
            ],
            "no_iep_explanation": [
                "If No, please explain",
            ],
            "aba_provider_obtained_iep": [
                "If yes, did the ABA provider obtain the current IEP/equivalent?",
            ],
            "current_iep_date": [
                "Date of the current IEP/equivalent",
            ],
            "aba_provider_participated_in_iep": [
                "Did the ABA provider participate in the IEP/equivalent meeting(s)?",
            ],
            "iep_attendance_dates": [
                "If yes, please indicate the date(s) of attendance",
            ],
            "iep_states_need_for_aba": [
                "Does the IEP/equivalent state the need for BHT/ABA services?",
            ],
        }

        school_iep = {}

        # Extract normal text values first
        for field_name, aliases in field_labels.items():
            value = None

            for row in table:
                if len(row) < 2:
                    continue

                label_cell = self.normalize_label(row[0])

                for alias in aliases:
                    normalized_alias = self.normalize_label(alias)

                    if (
                        label_cell == normalized_alias
                        or label_cell.startswith(normalized_alias + ":")
                    ):
                        value = row[1].strip()
                        break

                if value is not None:
                    break

            if value is not None:
                normalized_value = self.normalize_label(value)

                if any(
                    normalized_value == self.normalize_label(alias)
                    for alias in aliases
                ):
                    value = None

            if field_name == "no_iep_explanation" and value:
                prefix = "If No, please explain:"
                if value.lower().startswith(prefix.lower()):
                    value = value[len(prefix):].strip()
            school_iep[field_name] = value

        
        # Replaced Yes/No text with dynamically located Word checkboxes
        if self.docx_extractor is not None:
            school_iep["aba_services_requested_at_school"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Are ABA services being requested for authorization from CalOptima Health at the school setting?"
                ])
            )

            school_iep["has_current_iep"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Does the member have a current Individualized Educational Plan (IEP/equivalent)?"
                ])
            )

            school_iep["aba_provider_obtained_iep"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "If yes, did the ABA provider obtain the current IEP/equivalent?"
                ])
            )

            school_iep["aba_provider_participated_in_iep"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Did the ABA provider participate in the IEP/equivalent meeting(s)?"
                ])
            )

            school_iep["iep_states_need_for_aba"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Does the IEP/equivalent state the need for BHT/ABA services?"
                ])
            )

        return school_iep

    def extract_iep_services(self):
        table = self.find_table([
            "Individualized Educational Plan (IEP/equivalent) Information",
            "Service Type",
        ])

        if table is None:
            return []

        services = []

        expected_headers = {
            "service type",
            "location / name of school",
            "classroom type",
            "start date",
            "end date",
            "frequency",
        }

        header_row_index = None

        # Dynamically find the column-header row
        for row_index, row in enumerate(table):
            normalized_row = {
                self.normalize_label(cell)
                for cell in row
                if cell.strip()
            }

            if "service type" in normalized_row:
                header_row_index = row_index
                break

        if header_row_index is None:
            return []

        headers = table[header_row_index]

        # Map each column based on its header instead of fixed positions
        column_map = {}

        for column_index, header in enumerate(headers):
            normalized_header = self.normalize_label(header)

            if normalized_header in expected_headers:
                column_map[normalized_header] = column_index

        # Process all rows underneath the header
        for row in table[header_row_index + 1:]:

            values = {}

            for header, column_index in column_map.items():
                if column_index < len(row):
                    values[header] = row[column_index].strip()
                else:
                    values[header] = ""

            # Ignore completely empty service rows
            if not any(values.values()):
                continue

            services.append({
                "service_type": values.get("service type", ""),
                "location_school": values.get(
                    "location / name of school", ""
                ),
                "classroom_type": values.get("classroom type", ""),
                "start_date": values.get("start date", ""),
                "end_date": values.get("end date", ""),
                "frequency": values.get("frequency", ""),
            })

        return services


    def extract_previous_interventions(self):
        table = self.find_table([
            "PREVIOUS INTERVENTIONS",
            "Name of Provider",
        ])

        if table is None:
            return []

        expected_headers = {
            "name of provider",
            "service provided",
            "service level",
            "start date",
            "end date",
            "reason for termination",
        }

        header_row_index = None

        # Dynamically find the header row
        for row_index, row in enumerate(table):
            normalized_row = {
                self.normalize_label(cell)
                for cell in row
                if cell.strip()
            }

            if "name of provider" in normalized_row:
                header_row_index = row_index
                break

        if header_row_index is None:
            return []

        # Dynamically map column names to their positions
        column_map = {}

        for column_index, header in enumerate(table[header_row_index]):
            normalized_header = self.normalize_label(header)

            if normalized_header in expected_headers:
                column_map[normalized_header] = column_index

        interventions = []

        # Extract every populated row after the headers
        for row in table[header_row_index + 1:]:
            values = {}

            for header, column_index in column_map.items():
                if column_index < len(row):
                    values[header] = row[column_index].strip()
                else:
                    values[header] = ""

            if not any(values.values()):
                continue

            interventions.append({
                "provider_name": values.get("name of provider", ""),
                "service_provided": values.get("service provided", ""),
                "service_level": values.get("service level", ""),
                "start_date": values.get("start date", ""),
                "end_date": values.get("end date", ""),
                "reason_for_termination": values.get(
                    "reason for termination", ""
                ),
            })

        return interventions

    def extract_coordination_of_care(self):
        table = self.find_table([
            "Parent/Caregiver",
            "Regional Center",
            "Mental Health Provider",
        ])

        if table is None:
            return {}

        field_labels = {
            "parent_caregiver": [
                "Parent/Caregiver",
                "Parent / Caregiver",
            ],
            "school": [
                "School",
            ],
            "regional_center": [
                "Regional Center",
            ],
            "speech_ot_pt": [
                "Speech/OT/PT",
                "Speech / OT / PT",
            ],
            "primary_care_provider_specialist": [
                "Primary Care Provider/Specialist",
                "Primary Care Provider / Specialist",
                "PCP/Specialist",
            ],
            "mental_health_provider": [
                "Mental Health Provider",
            ],
        }

        coordination = {}

        for field_name, aliases in field_labels.items():
            value = self.find_value_after_labels(
                table,
                aliases,
                [
                    label
                    for alias_list in field_labels.values()
                    for label in alias_list
                ]
            )

            coordination[field_name] = value

        return coordination


    def extract_adaptive_testing(self):
        table = self.find_table([
            "Vineland-3 Scoring Information",
            "Date Administered",
            "Facilitated By",
        ])

        if table is None:
            return {}

        result = {
            "assessment": "Vineland-3",
            "baseline": {},
            "current": {},
            "scores": [],
        }

        baseline_columns = []
        current_columns = []
        score_header_row = None

        # Find Baseline / Current column groups and score header row
        for row_index, row in enumerate(table):
            normalized_row = [
                self.normalize_label(cell)
                for cell in row
            ]

            if "baseline" in normalized_row or "current" in normalized_row:
                baseline_columns = [
                    index
                    for index, value in enumerate(normalized_row)
                    if value == "baseline"
                ]

                current_columns = [
                    index
                    for index, value in enumerate(normalized_row)
                    if value == "current"
                ]

            if "domain" in normalized_row:
                score_header_row = row_index
                break

        if score_header_row is None:
            return result

        # Assessment metadata
        for row in table[:score_header_row]:
            if not row:
                continue

            label = self.normalize_label(row[0])

            if label == "date administered":
                if baseline_columns:
                    result["baseline"]["date_administered"] = (
                        row[baseline_columns[0]].strip()
                    )

                if current_columns:
                    result["current"]["date_administered"] = (
                        row[current_columns[0]].strip()
                    )

            elif label == "facilitated by":
                if baseline_columns:
                    result["baseline"]["facilitated_by"] = (
                        row[baseline_columns[0]].strip()
                    )

                if current_columns:
                    result["current"]["facilitated_by"] = (
                        row[current_columns[0]].strip()
                    )

            elif label == "respondent":
                if baseline_columns:
                    result["baseline"]["respondent"] = (
                        row[baseline_columns[0]].strip()
                    )

                if current_columns:
                    result["current"]["respondent"] = (
                        row[current_columns[0]].strip()
                    )

        # Score rows
        current_domain = None

        for row in table[score_header_row + 1:]:
            if not row:
                continue

            name = row[0].strip()

            if not name:
                continue

            remaining_values = [
                cell.strip()
                for cell in row[1:]
            ]

            # Domain heading rows have no scores
            if not any(remaining_values):
                current_domain = name
                continue

            baseline_raw = row[1].strip() if len(row) > 1 else ""
            baseline_standard = row[2].strip() if len(row) > 2 else ""
            baseline_age = row[3].strip() if len(row) > 3 else ""

            current_raw = row[4].strip() if len(row) > 4 else ""
            current_standard = row[5].strip() if len(row) > 5 else ""
            current_age = row[6].strip() if len(row) > 6 else ""

            result["scores"].append({
                "domain": current_domain,
                "skill_area": name,
                "baseline": {
                    "raw_score": baseline_raw,
                    "standard_v_scale_score": baseline_standard,
                    "age_equivalent": baseline_age,
                },
                "current": {
                    "raw_score": current_raw,
                    "standard_v_scale_score": current_standard,
                    "age_equivalent": current_age,
                },
            })

        return result

    def extract_diagnostic_information(self):
        table = self.find_table([
            "DIAGNOSTIC INFORMATION",
            "Current diagnosis code",
            "Diagnosis description",
        ])

        if table is None:
            return []

        header_aliases = {
            "diagnosis_code": [
                "Current diagnosis code",
                "Diagnosis code",
                "ICD code",
            ],
            "diagnosis_description": [
                "Diagnosis description",
                "Diagnosis",
            ],
            "diagnosis_date": [
                "Date of diagnosis/report",
                "Date of diagnosis",
                "Diagnosis date",
            ],
            "diagnosed_by": [
                "Diagnosed by (Full Name & credential)",
                "Diagnosed by",
            ],
        }

        header_row_index = None
        column_map = {}

        # Dynamically locate the header row
        for row_index, row in enumerate(table):
            temporary_map = {}

            for column_index, cell in enumerate(row):
                normalized_cell = self.normalize_label(cell)

                for field_name, aliases in header_aliases.items():
                    normalized_aliases = {
                        self.normalize_label(alias)
                        for alias in aliases
                    }

                    if normalized_cell in normalized_aliases:
                        temporary_map[field_name] = column_index
                        break

            # Require at least the diagnosis code and description
            if (
                "diagnosis_code" in temporary_map
                and "diagnosis_description" in temporary_map
            ):
                header_row_index = row_index
                column_map = temporary_map
                break

        if header_row_index is None:
            return []

        diagnoses = []

        # Extract all diagnosis rows below the headers
        for row in table[header_row_index + 1:]:
            diagnosis = {}

            for field_name, column_index in column_map.items():
                if column_index < len(row):
                    diagnosis[field_name] = row[column_index].strip()
                else:
                    diagnosis[field_name] = ""

            # Ignore completely empty rows
            if not any(diagnosis.values()):
                continue

            diagnoses.append(diagnosis)

        return diagnoses

    
    def extract_functional_assessments(self):
        field_labels = {
            "operational_definition": ["Operational Definition"],
            "course_of_behavior": ["Course of Behavior"],
            "baseline_data": ["Baseline Data"],
            "history_problem_chief_complaint": [
                "History of the Problem/Chief Complaint",
                "History of the Problem",
            ],
            "setting_events": ["Setting Events"],
            "trigger_events": ["Trigger Events"],
            "consequent_analysis": ["Consequent Analysis"],
            "hypothesized_function": [
                "Impressions and Analysis of Hypothesized Function",
                "Hypothesized Function",
            ],
        }

        assessments = []
        assessments_by_behavior = {}

        def first_value(row):
            return next(
                (cell.strip() for cell in row if cell and cell.strip()),
                "",
            )

        def match_field(text):
            normalized = self.normalize_label(text)

            for field_name, aliases in field_labels.items():
                for alias in aliases:
                    normalized_alias = self.normalize_label(alias)

                    if (
                        normalized == normalized_alias
                        or normalized.startswith(normalized_alias + ":")
                        or normalized.startswith(normalized_alias + " ")
                    ):
                        return field_name, alias

            return None, None

        def is_target_row(row):
            if not row:
                return False

            normalized = self.normalize_label(first_value(row))

            return (
                normalized == "target behavior"
                or normalized.startswith("target behavior ")
                or normalized.startswith("target behavior:")
            )

        def behavior_name_from_row(row):
            if len(row) > 1 and row[1] and row[1].strip():
                return row[1].strip()

            text = first_value(row)

            if ":" in text:
                return text.split(":", 1)[1].strip()

            return ""

        def get_field_value(table, row_index, alias):
            text = first_value(table[row_index])
            normalized_alias = self.normalize_label(alias)

            # Handle inline values:
            # Operational Definition: Description
            if ":" in text and "\n" not in text:
                prefix, remainder = text.split(":", 1)

                if self.normalize_label(prefix) == normalized_alias:
                    if remainder.strip():
                        return remainder.strip()

            # Handle values in the following row.
            if row_index + 1 >= len(table):
                return ""

            next_row = table[row_index + 1]
            next_value = first_value(next_row)

            if not next_value or is_target_row(next_row):
                return ""

            next_field, _ = match_field(next_value)

            if next_field is not None:
                return ""

            if self.normalize_label(next_value) in {
                "identifying behavior",
                "antecedent analysis",
                "behavior intervention plan",
            }:
                return ""

            return next_value

        # Search all tables, regardless of their headings.
        for table in self.tables:
            if not table:
                continue

            target_indices = [
                index
                for index, row in enumerate(table)
                if is_target_row(row)
            ]

            for position, start_index in enumerate(target_indices):
                end_index = (
                    target_indices[position + 1]
                    if position + 1 < len(target_indices)
                    else len(table)
                )

                # Identify whether this target-behavior block
                # contains a Functional Assessment.
                fields_found = set()

                for row in table[start_index + 1:end_index]:
                    field_name, _ = match_field(first_value(row))

                    if field_name:
                        fields_found.add(field_name)

                # Require assessment-specific labels to avoid
                # confusing BIP sections with assessments.
                required_fields = {
                    "operational_definition",
                    "course_of_behavior",
                    "baseline_data",
                }

                if not required_fields.issubset(fields_found):
                    continue

                behavior_name = behavior_name_from_row(
                    table[start_index]
                )

                if not behavior_name:
                    behavior_name = "Unspecified Target Behavior"

                behavior_key = behavior_name.casefold()

                if behavior_key not in assessments_by_behavior:
                    current_assessment = {
                        "target_behavior": behavior_name,
                    }

                    assessments_by_behavior[behavior_key] = (
                        current_assessment
                    )
                    assessments.append(current_assessment)

                else:
                    current_assessment = (
                        assessments_by_behavior[behavior_key]
                    )

                # Extract all assessment fields.
                for row_index in range(start_index + 1, end_index):
                    text = first_value(table[row_index])

                    field_name, alias = match_field(text)

                    if field_name is None:
                        continue

                    value = get_field_value(
                        table,
                        row_index,
                        alias,
                    )

                    if value:
                        current_assessment[field_name] = value

        return assessments


    
    def extract_behavior_intervention_plans(self):
        field_labels = {
            "ecological_interventions": [
                "Ecological interventions",
            ],
            "replacement_behavior_training": [
                "Training of replacement behaviors",
                "Replacement behavior training",
            ],
            "focused_intervention_strategies": [
                "Focused intervention strategies",
            ],
            "reactive_strategies": [
                "Reactive Strategies",
            ],
            "data_collection_procedures": [
                "Data collection procedures",
            ],
        }

        plans = []
        plans_by_behavior = {}

        def get_field_name(text):
            normalized = self.normalize_label(text)

            for field_name, aliases in field_labels.items():
                for alias in aliases:
                    normalized_alias = self.normalize_label(alias)

                    if (
                        normalized == normalized_alias
                        or normalized.startswith(normalized_alias + ":")
                        or normalized.startswith(normalized_alias + " ")
                    ):
                        return field_name, alias

            return None, None

        def get_first_value(row):
            for cell in row:
                if cell and cell.strip():
                    return cell.strip()
            return ""

        def extract_field_value(table, row_index, field_name, alias):
            row = table[row_index]
            first_cell = get_first_value(row)

            normalized_alias = self.normalize_label(alias)
            normalized = self.normalize_label(first_cell)

            value = ""

            # Determine whether the row contains a heading
            # or a heading followed by actual content.
            #
            # Template headings with multiline instructions
            # are treated as labels, not extracted values.

            is_label_row = (
                normalized == normalized_alias
                or "\n" in first_cell
            )

            if not is_label_row and ":" in first_cell:
                # Extract content following an inline label.
                prefix, remainder = first_cell.split(":", 1)

                if self.normalize_label(prefix) == normalized_alias:
                    value = remainder.strip()

            # If the label row contains no inline value,
            # retrieve the content from the next row.
            if not value and row_index + 1 < len(table):
                next_value = get_first_value(table[row_index + 1])

                next_normalized = self.normalize_label(next_value)
                next_field, _ = get_field_name(next_value)

                # Avoid extracting another section heading
                # or target behavior as a field value.
                if (
                    next_value
                    and not next_normalized.startswith("target behavior")
                ):
                    if next_field is None:
                        value = next_value

                    elif (
                        next_field == field_name
                        and ":" in next_value
                        and "\n" not in next_value
                    ):
                        # Some templates repeat the label in
                        # the value row:
                        #
                        # Ecological interventions: Description
                        #
                        # Extract only the description.
                        prefix, remainder = next_value.split(":", 1)

                        if (
                            self.normalize_label(prefix)
                            == normalized_alias
                        ):
                            value = remainder.strip()

            return value

        for table in self.tables:
            if not table:
                continue

            # Identify Behavior Intervention Plan tables.
            heading = " ".join(
                cell.strip()
                for row in table[:2]
                for cell in row
                if cell and cell.strip()
            ).lower()

            if "behavior intervention plan" not in heading:
                continue

            current_plan = None

            for row_index, row in enumerate(table):
                if not row:
                    continue

                first_cell = get_first_value(row)
                normalized = self.normalize_label(first_cell)

                # ------------------------------------
                # TARGET BEHAVIOR
                # ------------------------------------
                if normalized.startswith("target behavior"):
                    behavior_name = ""

                    if len(row) > 1 and row[1]:
                        behavior_name = row[1].strip()

                    if not behavior_name:
                        behavior_name = "Unspecified Target Behavior"

                    behavior_key = behavior_name.casefold()

                    if behavior_key not in plans_by_behavior:
                        current_plan = {
                            "target_behavior": behavior_name,
                        }

                        plans_by_behavior[behavior_key] = current_plan
                        plans.append(current_plan)

                    else:
                        current_plan = plans_by_behavior[behavior_key]

                    continue

                if current_plan is None:
                    continue

                # ------------------------------------
                # INTERVENTION FIELDS
                # ------------------------------------
                field_name, alias = get_field_name(first_cell)

                if field_name is None:
                    continue

                value = extract_field_value(
                    table,
                    row_index,
                    field_name,
                    alias,
                )

                if value:
                    current_plan[field_name] = value

        # Remove entries without any intervention data.
        plans = [
            plan
            for plan in plans
            if any(
                field_name in plan
                for field_name in field_labels
            )
        ]

        return plans



    def extract_mediator_analysis(self):
        table = self.find_table([
            "MEDIATOR ANALYSIS",
            "Mediator Analysis",
        ])

        if table is None:
            return {}

        # Find the Mediator Analysis heading dynamically
        for row_index, row in enumerate(table):
            for cell in row:
                if self.normalize_label(cell) == "mediator analysis":

                    # Search the rows after the heading.
                    # Skip the instruction row and return the
                    # first populated narrative row.
                    for next_row in table[row_index + 1:]:
                        non_empty_values = [
                            value.strip()
                            for value in next_row
                            if value.strip()
                        ]

                        if not non_empty_values:
                            continue

                        value = non_empty_values[0]

                        # Skip the template instruction
                        if value.lower().startswith(
                            "provide additional information"
                        ):
                            continue

                        return {
                            "mediator_analysis": value
                        }

        return {}

    def extract_reinforcer_assessment(self):
        table = self.find_table([
            "REINFORCER ASSESSMENT",
            "Reinforcer Assessment",
        ])

        if table is None:
            return {}

        # Dynamically locate the Reinforcer Assessment heading
        for row_index, row in enumerate(table):
            for cell in row:
                if self.normalize_label(cell) == "reinforcer assessment":

                    # Search all rows following the heading
                    for next_row in table[row_index + 1:]:
                        non_empty_values = [
                            value.strip()
                            for value in next_row
                            if value.strip()
                        ]

                        if not non_empty_values:
                            continue

                        value = non_empty_values[0]

                        # Skip template instruction text
                        if value.lower().startswith(
                            "describe reinforcers identified"
                        ):
                            continue

                        return {
                            "reinforcer_assessment": value
                        }

        return {}

    def extract_target_replacement_goals(self):
        target_goals = []
        replacement_goals = []

        # Keep current_goal outside the table loop so a goal
        # can continue into the next Word table.
        current_goal = None

        for table in self.tables:

            for row_index, row in enumerate(table):
                if not row:
                    continue

                first_value = next(
                    (
                        cell.strip()
                        for cell in row
                        if cell.strip()
                    ),
                    "",
                )

                if not first_value:
                    continue

                normalized_value = self.normalize_label(first_value)

                # -------------------------
                # TARGET BEHAVIOR GOAL
                # -------------------------
                if normalized_value.startswith("target behavior goal"):
                    current_goal = {
                        "goal": first_value,
                        "topography": "",
                        "location_setting": "",
                        "antecedent_strategies": "",
                        "consequent_strategies": "",
                        "date_of_introduction": "",
                        "baseline_data_date": "",
                    }

                    target_goals.append(current_goal)
                    continue

                # -------------------------
                # REPLACEMENT BEHAVIOR GOAL
                # -------------------------
                if normalized_value.startswith("replacement behavior goal"):
                    current_goal = {
                        "goal": first_value,
                        "date_of_introduction": "",
                        "location_setting": "",
                        "baseline_data_date": "",
                    }

                    replacement_goals.append(current_goal)
                    continue

                if current_goal is None:
                    continue

                # -------------------------
                # TARGET GOAL FIELDS
                # -------------------------
                if normalized_value.startswith(
                    "topography of target behavior"
                ):
                    current_goal["topography"] = first_value

                elif normalized_value.startswith(
                    "antecedent strategies for target behavior"
                ):
                    current_goal["antecedent_strategies"] = first_value

                elif normalized_value.startswith(
                    "consequent strategies for target behavior"
                ):
                    current_goal["consequent_strategies"] = first_value

                # -------------------------
                # DATE OF INTRODUCTION
                # -------------------------
                elif normalized_value.startswith("date of introduction"):
                    current_goal["date_of_introduction"] = first_value

                # Handles things like:
                # "1.Date of Introduction: August 2025"
                elif (
                    "date of introduction" in normalized_value
                    and normalized_value[0].isdigit()
                ):
                    current_goal["date_of_introduction"] = first_value

                # -------------------------
                # BASELINE
                # -------------------------
                elif normalized_value.startswith(
                    "baseline data and date"
                ):
                    current_goal["baseline_data_date"] = first_value

                # -------------------------
                # LOCATION / SETTING
                # -------------------------
                elif (
                    normalized_value.startswith(
                        "location/setting of target behavior"
                    )
                    or normalized_value.startswith(
                        "location/setting of replacement behavior"
                    )
                ):
                    location = ""

                    # Location choices are in the following row
                    if row_index + 1 < len(table):
                        next_row = table[row_index + 1]

                        location = next(
                            (
                                cell.strip()
                                for cell in next_row
                                if cell.strip()
                            ),
                            "",
                        )

                    current_goal["location_setting"] = location

        return {
            "target_behavior_goals": target_goals,
            "replacement_behavior_goals": replacement_goals,
        }

    
    def extract_skill_acquisition_goals(self):
        import re

        intervention_areas = []
        current_area = None
        current_goal = None
        expecting_description = False

        def extract_after_colon(value):
            return value.split(":", 1)[1].strip() if ":" in value else value.strip()

        for table in self.tables:
            for row in table:
                if not row:
                    continue

                first_value = next(
                    (cell.strip() for cell in row if cell.strip()),
                    ""
                )

                if not first_value:
                    continue

                normalized_value = self.normalize_label(first_value)

                # Stop when we reach a different major section
                if current_area is not None and (
                    normalized_value.startswith("parent/caregiver goals")
                    or normalized_value.startswith("parent/caregiver goal")
                    or normalized_value.startswith("plan for generalization")
                    or normalized_value.startswith("transition plan")
                    or normalized_value.startswith("crisis plan")
                    or normalized_value.startswith("summary and recommendations")
                ):
                    current_area = None
                    current_goal = None
                    expecting_description = False
                    continue

                # Intervention area
                if normalized_value.startswith("intervention area"):
                    current_area = {
                        "intervention_area": extract_after_colon(first_value),
                        "goals": []
                    }
                    intervention_areas.append(current_area)
                    current_goal = None
                    expecting_description = False
                    continue

                if current_area is None:
                    continue

                # Goal heading (optional numbering before the heading)
                goal_match = re.match(
                    r"^(?:\d+\.\s*)?skill acquisition goal\s*\d+\s*:",
                    first_value,
                    re.IGNORECASE
                )

                if goal_match:
                    current_goal = {
                        "goal_label": goal_match.group(0).strip(),
                        "goal_description": "",
                        "date_of_introduction": "",
                        "location_setting": "",
                        "baseline_data_date": ""
                    }

                    current_area["goals"].append(current_goal)
                    expecting_description = True

                    # Handle a description in the same cell
                    remaining_text = first_value[goal_match.end():].strip()
                    if remaining_text:
                        current_goal["goal_description"] = remaining_text
                        expecting_description = False

                    continue

                if current_goal is None:
                    continue

                # Date of introduction
                if normalized_value.startswith("date of introduction"):
                    current_goal["date_of_introduction"] = extract_after_colon(
                        first_value
                    )
                    expecting_description = False
                    continue

                # Location / setting
                if normalized_value.startswith(
                    "location/setting of skill acquisition goal"
                ):
                    current_goal["location_setting"] = extract_after_colon(
                        first_value
                    )
                    expecting_description = False
                    continue

                # Baseline data
                if (
                    normalized_value.startswith("baseline data with dates")
                    or normalized_value.startswith("baseline data and date")
                ):
                    current_goal["baseline_data_date"] = extract_after_colon(
                        first_value
                    )
                    expecting_description = False
                    continue

                # Description immediately following goal heading
                if expecting_description:
                    current_goal["goal_description"] = first_value
                    expecting_description = False

        return intervention_areas


    
    
    def extract_parent_caregiver_goals(self):
        participants = []
        goals = []

        in_section = False
        in_goals = False
        in_participants = False
        current_goal = None

        end_section_headings = {
            "generalization and maintenance plan",
            "transition plan",
            "behavior intervention plan",
        }

        for table in self.tables:
            for row in table:
                if not row:
                    continue

                values = [
                    cell.strip()
                    for cell in row
                    if cell and cell.strip()
                ]

                if not values:
                    continue

                first_value = values[0]
                normalized = self.normalize_label(first_value)

                # -------------------------
                # START SECTION
                # -------------------------
                if normalized == "parent/caregiver goals":
                    in_section = True
                    in_goals = False
                    in_participants = False
                    current_goal = None
                    continue

                if not in_section:
                    continue

                # -------------------------
                # END SECTION
                # -------------------------
                if (
                    normalized.startswith("plan for generalization")
                    or normalized in end_section_headings
                ):
                    in_section = False
                    in_goals = False
                    in_participants = False
                    current_goal = None
                    continue

                # -------------------------
                # PARTICIPANT TABLE HEADER
                # -------------------------
                if normalized.startswith(
                    "participants in parent training"
                ):
                    in_participants = True
                    in_goals = False
                    continue

                # -------------------------
                # SKIP INSTRUCTIONS
                # -------------------------
                if (
                    "parent/caregiver goals should specify" in normalized
                    or normalized.startswith("previously mastered skills")
                    or normalized.startswith("identify measurable")
                ):
                    continue

                # -------------------------
                # GOAL START (LABELED)
                # -------------------------
                if (
                    normalized.startswith("parent/caregiver goal")
                    and "long-term" in normalized
                ):
                    in_participants = False
                    in_goals = True

                    current_goal = {
                        "goal": first_value,
                        "date_initiated": "",
                        "location_setting": "",
                        "baseline_data_date": "",
                    }

                    goals.append(current_goal)
                    continue

                # -------------------------
                # GOAL START (UNLABELED)
                # -------------------------
                if normalized.startswith("by "):
                    in_participants = False
                    in_goals = True

                    current_goal = {
                        "goal": first_value,
                        "date_initiated": "",
                        "location_setting": "",
                        "baseline_data_date": "",
                    }

                    goals.append(current_goal)
                    continue

                # -------------------------
                # END PARTICIPANT TABLE
                # -------------------------
                if (
                    in_participants
                    and normalized.startswith("parent/caregiver goal")
                ):
                    in_participants = False
                    continue

                # -------------------------
                # PARTICIPANTS
                # -------------------------
                if in_participants and not in_goals:
                    participants.append({
                        "participant": first_value,
                        "previously_mastered_skills": (
                            row[1].strip()
                            if len(row) > 1 and row[1]
                            else ""
                        ),
                    })
                    continue

                if current_goal is None:
                    continue

                # -------------------------
                # DATE INITIATED
                # -------------------------
                if normalized.startswith("date initiated"):
                    current_goal["date_initiated"] = first_value
                    continue

                # -------------------------
                # BASELINE DATA
                # -------------------------
                if (
                    normalized.startswith("baseline data and date")
                    or normalized.startswith("baseline data with dates")
                ):
                    current_goal["baseline_data_date"] = first_value
                    continue

                # -------------------------
                # LOCATION / SETTING
                # -------------------------
                if (
                    normalized.startswith("parent/caregiver goal")
                    and "long-term" not in normalized
                    and any(
                        word in normalized
                        for word in (
                            "telehealth",
                            "home",
                            "school",
                            "clinic",
                            "community",
                        )
                    )
                ):
                    current_goal["location_setting"] = (
                        first_value.split(":", 1)[1].strip()
                        if ":" in first_value
                        else first_value
                    )
                    continue

        return {
            "participants": participants,
            "goals": goals,
        }



    def extract_generalization_maintenance_plan(self):
        field_labels = {
            "data_collection": [
                "Data Collection",
            ],
            "procedural_reliability": [
                "Procedural Reliability",
            ],
            "reinforcement_schedule_thinning": [
                "Thinning the Schedule of Reinforcement",
            ],
            "family_caregiver_training_monitoring": [
                "Family/Caregiver Training and Monitoring",
            ],
            "transition_to_natural_mediators": [
                "Transition to Natural Mediators",
            ],
            "relapse_prevention": [
                "Relapse Prevention",
            ],
            "generalization_maintenance_plan": [
                "Generalization and Maintenance Plan",
            ],
        }

        plan = {}

        def first_value(row):
            return next(
                (
                    cell.strip()
                    for cell in row
                    if cell and cell.strip()
                ),
                "",
            )

        def match_field(text):
            for field_name, aliases in field_labels.items():
                for alias in aliases:
                    # The label may contain optional instructions:
                    # Data Collection (include plan to address):
                    #
                    # Match the beginning of the label while
                    # requiring a label boundary afterward.

                    normalized_text = text.strip().casefold()
                    normalized_alias = alias.casefold()

                    if normalized_text.startswith(normalized_alias):
                        remainder = normalized_text[
                            len(normalized_alias):
                        ]

                        if (
                            not remainder
                            or remainder.startswith(":")
                            or remainder.startswith(" (")
                        ):
                            return field_name

            return None

        for table in self.tables:
            if not table:
                continue

            # Identify the actual Generalization &
            # Maintenance Plan section by its heading.
            heading = first_value(table[0]).casefold()

            if not (
                "plan for generalization" in heading
                and "maintenance" in heading
            ):
                continue

            for row_index, row in enumerate(table[1:], start=1):
                text = first_value(row)

                if not text:
                    continue

                field_name = match_field(text)

                if field_name is None:
                    continue

                value = ""

                # CASE 1:
                # Label and description are in the same row.
                #
                # Data Collection (include plan to address):
                # Description...

                if ":" in text:
                    value = text.split(":", 1)[1].strip()

                # CASE 2:
                # Label is in one row, value is in the next.
                #
                # Data Collection:
                # Description...

                if not value:
                    for next_index in range(
                        row_index + 1,
                        len(table),
                    ):
                        next_text = first_value(
                            table[next_index]
                        )

                        if not next_text:
                            continue

                        # Don't mistake the next field label
                        # for a description.
                        if match_field(next_text) is not None:
                            break

                        value = next_text
                        break

                if value:
                    plan[field_name] = value

        return plan

    
    def extract_transition_plan(self):
        table = self.find_table([
            "TRANSITION PLAN",
            "Transition Plan",
        ])

        if table is None:
            return {}

        result = {
            "treatment_report_status": None,
            "treatment_report_selection_method": None,
            "exit_plan_criteria": "",
            "individualized_timeline": "",
            "generalization_expectations": "",
            "service_fading_plan": "",
        }

        question_labels = {
            "please list exit plan/criteria":
                "exit_plan_criteria",
            "please include an individualized timeline":
                "individualized_timeline",
            "what is the expectation for generalization":
                "generalization_expectations",
            "please provide a description of how the level of services will":
                "service_fading_plan",
        }

        treatment_options = {
            "final treatment report":
                "Final treatment report",
            "treatment report transition to parent/caregiver":
                "Transition to parent/caregiver or social skills group",
            "treatment report with fading":
                "Fading of interventionist direct hours",
            "treatment report initial/ongoing relationship":
                "Initial/ongoing ABA services",
        }

        # -----------------------------------
        # EXTRACT NARRATIVE FIELDS
        # -----------------------------------
        for row_index, row in enumerate(table):
            if not row:
                continue

            first_value = next(
                (cell.strip() for cell in row if cell.strip()),
                ""
            )

            if not first_value:
                continue

            normalized_value = self.normalize_label(first_value)

            for label, field_name in question_labels.items():
                if normalized_value.startswith(label):

                    for next_row in table[row_index + 1:]:
                        answer = next(
                            (cell.strip() for cell in next_row if cell.strip()),
                            ""
                        )

                        if answer:
                            result[field_name] = answer
                            break

                    break

        # -----------------------------------
        # DETECT TREATMENT REPORT SELECTION
        # -----------------------------------
        # The original narrative extraction uses self.tables.
        # For ink drawings, we also inspect the underlying DOCX XML.

        from docx import Document

        if self.docx_extractor is None:
            return result

        doc = self.docx_extractor.document

        for doc_table in doc.tables:
            table_text = " ".join(
                cell.text
                for row in doc_table.rows
                for cell in row.cells
            )

            if "TRANSITION PLAN" not in table_text.upper():
                continue

            for row in doc_table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        option_text = paragraph.text.strip()
                        normalized_option = self.normalize_label(option_text)

                        for prefix, option_name in treatment_options.items():
                            if not normalized_option.startswith(prefix):
                                continue

                            # Standard Word checkbox content controls
                            checkbox_elements = paragraph._p.xpath(
                                ".//*[local-name()='checkbox']"
                            )

                            checked = False
                            for checkbox in checkbox_elements:
                                states = checkbox.xpath(
                                    ".//*[local-name()='checked']"
                                )
                                for state in states:
                                    value = state.get(
                                        "{http://schemas.microsoft.com/office/word/2010/wordml}val"
                                    )
                                    if value in ("1", "true", "on"):
                                        checked = True

                            if checked:
                                result["treatment_report_status"] = option_name
                                result["treatment_report_selection_method"] = (
                                    "word_checkbox"
                                )
                                return result

                            # Ink drawings can indicate handwritten selections.
                            ink_parts = paragraph._p.xpath(
                                ".//*[local-name()='contentPart']"
                            )

                            if ink_parts:
                                result["treatment_report_status"] = option_name
                                result["treatment_report_selection_method"] = (
                                    "ink_drawing_unverified"
                                )
                                return result

        return result


    def extract_crisis_plan(self):
        table = self.find_table([
            "CRISIS PLAN",
            "Crisis Plan",
        ])

        if table is None:
            return {}

        for row_index, row in enumerate(table):
            for cell in row:
                normalized_cell = self.normalize_label(cell)

                # Dynamically locate the Crisis Plan heading
                if normalized_cell == "crisis plan":

                    # Search everything following the heading
                    for next_row in table[row_index + 1:]:
                        value = next(
                            (
                                cell.strip()
                                for cell in next_row
                                if cell.strip()
                            ),
                            "",
                        )

                        if not value:
                            continue

                        normalized_value = self.normalize_label(value)

                        # Skip the template instruction
                        if normalized_value.startswith(
                            "define what steps the member"
                        ):
                            continue

                        return {
                            "crisis_plan": value
                        }

        return {}

    
    def extract_summary_recommendations(self):
        import re

        result = {
            "clinical_summary": "",
            "service_requests": [],
        }

        def unique_cells(row):
            return list(dict.fromkeys(
                cell.strip()
                for cell in row
                if cell and cell.strip()
            ))

        def is_instruction(text):
            normalized = self.normalize_label(text)

            instruction_patterns = [
                "provide a clinical summary",
                "consider the member",
                "justify hours requested",
                "justifies hours requested",
            ]

            return any(
                pattern in normalized
                for pattern in instruction_patterns
            )

        def is_hcpcs_header(row):
            combined = " ".join(
                cell.strip().lower()
                for cell in row
                if cell
            )

            return (
                "hcpcs" in combined
                and "description" in combined
            )

        def is_hcpcs_code(text):
            return bool(
                re.fullmatch(
                    r"[A-Z]\d{4}(?:-[A-Z0-9]+)*",
                    text.strip().upper(),
                )
            )

        # Search all tables rather than relying
        # on a fixed table number.
        for table in self.tables:
            if not table:
                continue

            heading = " ".join(
                unique_cells(table[0])
            ).casefold()

            if not (
                "summary and recommendations" in heading
                or "clinical summary" in heading
            ):
                continue

            # ----------------------------------
            # FIND SERVICE TABLE HEADER
            # ----------------------------------
            header_row_index = None
            header_map = {}

            for row_index, row in enumerate(table):
                if not is_hcpcs_header(row):
                    continue

                header_row_index = row_index

                for column_index, cell in enumerate(row):
                    normalized = self.normalize_label(cell)

                    if "hcpcs" in normalized and "code" in normalized:
                        header_map["hcpcs_code_modifiers"] = column_index

                    elif normalized == "description":
                        header_map["description"] = column_index

                    elif (
                        "total hours requested per month" in normalized
                        and "telehealth" not in normalized
                    ):
                        header_map["hours_per_month"] = column_index

                    elif "total units requested per 6 month" in normalized:
                        header_map["units_per_6_months"] = column_index

                    elif "location of service" in normalized:
                        header_map["location_of_service"] = column_index

                    elif "telehealth" in normalized and "hours" in normalized:
                        header_map["telehealth_hours_per_month"] = column_index

                break

            # ----------------------------------
            # EXTRACT CLINICAL SUMMARY
            # ----------------------------------
            narrative_end = (
                header_row_index
                if header_row_index is not None
                else len(table)
            )

            narrative_parts = []

            for row in table[1:narrative_end]:
                for text in unique_cells(row):
                    if is_instruction(text):
                        continue

                    narrative_parts.append(text)

            if narrative_parts:
                result["clinical_summary"] = "\n".join(
                    narrative_parts
                )

            # ----------------------------------
            # EXTRACT SERVICE REQUESTS
            # ----------------------------------
            if header_row_index is None:
                continue

            code_column = header_map.get(
                "hcpcs_code_modifiers"
            )

            if code_column is None:
                continue

            for row in table[header_row_index + 1:]:
                if code_column >= len(row):
                    continue

                code = row[code_column].strip()

                if not is_hcpcs_code(code):
                    continue

                service = {}

                for field_name, column_index in header_map.items():
                    if column_index < len(row):
                        service[field_name] = (
                            row[column_index].strip()
                            if row[column_index]
                            else ""
                        )
                    else:
                        service[field_name] = ""

                result["service_requests"].append(service)

        return result

    def extract_parent_guardian_involvement(self):
        result = {
            "involved_in_treatment_plan": None,
            "agrees_with_treatment_plan": None,
            "explanation_if_no": "",
        }

        # -------------------------
        # YES / NO CHECKBOXES
        # -------------------------
        if self.docx_extractor is not None:
            result["involved_in_treatment_plan"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Was the Parent/guardian involved in the development of the treatment plan?",
                    "Was the Parent/guardian involved in the development of the treatment plan",
                ])
            )

            result["agrees_with_treatment_plan"] = (
                self.docx_extractor.get_yes_no_checkbox_by_label([
                    "Is the parent/guardian in agreement with the submitted treatment plan?",
                    "Is the parent/guardian in agreement with the submitted treatment plan",
                ])
            )

        # -------------------------
        # EXPLANATION IF NO
        # -------------------------
        table = self.find_table([
            "PARENT/CAREGIVER OR LEGAL GUARDIAN INVOLVEMENT",
            "Parent/Caregiver or Legal Guardian Involvement",
        ])

        if table is None:
            return result

        explanation_labels = [
            "If No to any response, please provide an explanation",
        ]

        for row_index, row in enumerate(table):
            for column_index, cell in enumerate(row):
                normalized_cell = self.normalize_label(cell)

                if any(
                    normalized_cell.startswith(
                        self.normalize_label(label)
                    )
                    for label in explanation_labels
                ):
                    # First try another cell in the same row
                    for value_cell in row[column_index + 1:]:
                        value = value_cell.strip()

                        if (
                            value
                            and not value.lower().startswith(
                                "if no to any response"
                            )
                        ):
                            result["explanation_if_no"] = value
                            return result

                    # Otherwise search following rows
                    for next_row in table[row_index + 1:]:
                        value = next(
                            (
                                next_cell.strip()
                                for next_cell in next_row
                                if next_cell.strip()
                            ),
                            "",
                        )

                        if value:
                            result["explanation_if_no"] = value
                            return result

        return result
    
    
    
    
    def extract_signatures(self):
        from docx.oxml.ns import qn

        result = {
            "report_writer": {
                "name_credentials": "",
                "title_license": "",
                "report_completion_date": "",
                "signature_date": "",
                "signature_present": None,
            },
            "report_reviewer": {
                "name_credentials": "",
                "title_license": "",
                "report_completion_date": "",
                "signature_date": "",
                "signature_present": None,
            },
        }

        def get_value_after_label(row):
            if not row:
                return ""

            label = self.normalize_label(row[0])

            for cell in row[1:]:
                value = cell.strip() if cell else ""

                if not value:
                    continue

                if self.normalize_label(value) == label:
                    continue

                return value

            return ""

        def has_embedded_image(cell, document_part):
            """
            Check whether a Word table cell contains
            an embedded image relationship.
            Supports DrawingML and legacy VML images.
            """
            element = cell._element

            relationship_ids = set()

            # Modern Word images
            for blip in element.xpath(".//a:blip"):
                relationship_id = blip.get(qn("r:embed"))

                if relationship_id:
                    relationship_ids.add(relationship_id)

            # Legacy VML images
            for child in element.iter():
                if child.tag.endswith("}imagedata"):
                    relationship_id = child.get(qn("r:id"))

                    if relationship_id:
                        relationship_ids.add(relationship_id)

            for relationship_id in relationship_ids:
                relationship = document_part.rels.get(
                    relationship_id
                )

                if relationship is None:
                    continue

                if relationship.reltype.endswith("/image"):
                    return True

            return False

        document = self.docx_extractor.document

        # Use original DOCX tables for image detection.
        for table_index, table in enumerate(self.tables):
            if not table:
                continue

            heading = " ".join(
                cell.strip().lower()
                for cell in table[0]
                if cell
            )

            if "report written by" in heading:
                role = "report_writer"

            elif "report reviewed by" in heading:
                role = "report_reviewer"

            else:
                continue

            signature = result[role]

            # ----------------------------------
            # EXTRACT SIGNATURE METADATA
            # ----------------------------------
            for row in table:
                if not row:
                    continue

                first_cell = row[0].strip() if row[0] else ""
                normalized = self.normalize_label(first_cell)

                if (
                    "report written by" in normalized
                    or "report reviewed by" in normalized
                ):
                    signature["name_credentials"] = (
                        get_value_after_label(row)
                    )

                elif (
                    normalized.startswith("title")
                    and "license" in normalized
                ):
                    signature["title_license"] = (
                        get_value_after_label(row)
                    )

                elif "date of report completed" in normalized:
                    signature["report_completion_date"] = (
                        get_value_after_label(row)
                    )

                elif normalized.startswith("signature"):
                    for index, cell in enumerate(row):
                        if not cell:
                            continue

                        if self.normalize_label(cell) == "date":
                            for next_cell in row[index + 1:]:
                                if next_cell and next_cell.strip():
                                    signature["signature_date"] = (
                                        next_cell.strip()
                                    )
                                    break

                            break

            # ----------------------------------
            # DETECT EMBEDDED SIGNATURE IMAGE
            # ----------------------------------
            if table_index >= len(document.tables):
                continue

            original_table = document.tables[table_index]

            for row_index, row in enumerate(table):
                if not row:
                    continue

                first_cell = row[0].strip() if row[0] else ""

                if not self.normalize_label(
                    first_cell
                ).startswith("signature"):
                    continue

                original_row = original_table.rows[row_index]

                image_found = False

                for cell in original_row.cells:
                    if has_embedded_image(
                        cell,
                        document.part,
                    ):
                        image_found = True
                        break

                signature["signature_present"] = image_found

        return result





    def extract_telehealth_consent(self):
        result = {
            "telehealth_consent": None,
            "consent_date": "",
            "no_telehealth_services": None,
        }

        table = self.find_table([
            "Telehealth Consent Confirmation",
            "Telehealth Consent",
        ])

        if table is None:
            return result

        # Extract consent date dynamically
        for row in table:
            for cell in row:
                value = cell.strip()

                if self.normalize_label(value).startswith(
                    "if yes, please confirm the date consent obtained"
                ):
                    if ":" in value:
                        result["consent_date"] = value.split(
                            ":", 1
                        )[1].strip()

        # Extract checkbox states from Word XML
        if self.docx_extractor is not None:

            for word_table in self.docx_extractor.document.tables:

                heading_found = any(
                    "telehealth consent confirmation" in cell.text.lower()
                    for row in word_table.rows
                    for cell in row.cells
                )

                if not heading_found:
                    continue

                for row in word_table.rows:
                    for cell in row.cells:

                        if "consent was obtained" not in cell.text.lower():
                            continue

                        checkboxes = cell._tc.xpath(
                            './/*[local-name()="checkbox"]'
                        )

                        checkbox_states = []

                        for checkbox in checkboxes:
                            checked_elements = checkbox.xpath(
                                './*[local-name()="checked"]'
                            )

                            if not checked_elements:
                                checkbox_states.append(False)
                                continue

                            checked_element = checked_elements[0]

                            checked_value = next(
                                (
                                    attribute_value
                                    for attribute_name, attribute_value
                                    in checked_element.attrib.items()
                                    if attribute_name.endswith("}val")
                                    or attribute_name == "val"
                                ),
                                None
                            )

                            checkbox_states.append(
                                checked_value == "1"
                            )

                        if len(checkbox_states) == 3:
                            yes_checked = checkbox_states[0]
                            no_checked = checkbox_states[1]
                            na_checked = checkbox_states[2]

                            if sum(checkbox_states) == 1:
                                result["telehealth_consent"] = (
                                    "Yes" if yes_checked
                                    else "No" if no_checked
                                    else "N/A"
                                )
                                result["no_telehealth_services"] = na_checked

                        return result

        return result

                