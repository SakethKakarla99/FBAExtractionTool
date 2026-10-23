from docx import Document
from docx.oxml.ns import qn


class DOCXExtractor:
    def __init__(self, file_path):
        self.file_path = file_path
        self.document = Document(file_path)

    def get_paragraphs(self):
        paragraphs = []

        for paragraph in self.document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        return paragraphs
    def get_tables(self):
        tables = []

        for table in self.document.tables:
            table_data = []

            for row in table.rows:
                row_data = []

                for cell in row.cells:
                    text = cell.text.strip()
                    row_data.append(text)

                table_data.append(row_data)

            tables.append(table_data)

        return tables

    def get_yes_no_checkbox(self, table_index, row_index, cell_index):
        cell = self.document.tables[table_index].rows[row_index].cells[cell_index]

        checkbox_elements = cell._tc.xpath(
            './/*[local-name()="checkbox"]'
        )

        checkbox_states = []

        for checkbox in checkbox_elements:
            checked_elements = checkbox.xpath(
                './*[local-name()="checked"]'
            )

            if checked_elements:
                checked_element = checked_elements[0]

                value = None

                for attribute_name, attribute_value in checked_element.attrib.items():
                    if attribute_name.endswith("}val") or attribute_name == "val":
                        value = attribute_value
                        break

                checkbox_states.append(value == "1")

        if len(checkbox_states) < 2:
            return None

        yes_checked = checkbox_states[0]
        no_checked = checkbox_states[1]

        if yes_checked and not no_checked:
            return "Yes"

        if no_checked and not yes_checked:
            return "No"

        return None

    
    def get_yes_no_checkbox_by_label(self, possible_labels):
        import re

        normalized_labels = [
            " ".join(label.strip().lower().rstrip(":").split())
            for label in possible_labels
        ]

        for table in self.document.tables:
            for row in table.rows:
                for cell in row.cells:
                    normalized_cell = " ".join(
                        cell.text.strip().lower().rstrip(":").split()
                    )

                    label_found = any(
                        normalized_cell == label
                        or normalized_cell.startswith(label + ":")
                        for label in normalized_labels
                    )

                    if not label_found:
                        continue

                    for value_cell in row.cells:
                        # First: detect Unicode checkbox symbols
                        text = value_cell.text

                        yes_match = re.search(
                            r"([☒☑☐])\s*Yes\b",
                            text,
                            re.IGNORECASE
                        )

                        no_match = re.search(
                            r"([☒☑☐])\s*No\b",
                            text,
                            re.IGNORECASE
                        )

                        if yes_match and no_match:
                            yes_checked = yes_match.group(1) in ("☒", "☑")
                            no_checked = no_match.group(1) in ("☒", "☑")

                            if yes_checked and not no_checked:
                                return "Yes"

                            if no_checked and not yes_checked:
                                return "No"

                        # Second: preserve legacy Word checkbox support
                        checkbox_elements = value_cell._tc.xpath(
                            './/*[local-name()="checkbox"]'
                        )

                        if len(checkbox_elements) < 2:
                            continue

                        checkbox_states = []

                        for checkbox in checkbox_elements:
                            checked_elements = checkbox.xpath(
                                './*[local-name()="checked"]'
                            )

                            if not checked_elements:
                                continue

                            checked_element = checked_elements[0]
                            value = None

                            for attribute_name, attribute_value in checked_element.attrib.items():
                                if (
                                    attribute_name.endswith("}val")
                                    or attribute_name == "val"
                                ):
                                    value = attribute_value
                                    break

                            checkbox_states.append(
                                value in ("1", "true", "on")
                            )

                        if len(checkbox_states) < 2:
                            continue

                        yes_checked = checkbox_states[0]
                        no_checked = checkbox_states[1]

                        if yes_checked and not no_checked:
                            return "Yes"

                        if no_checked and not yes_checked:
                            return "No"

        return None


    def get_signature_presence_by_role(self, role_label):
        for table in self.document.tables:

            # Identify the signature table dynamically
            table_text = " ".join(
                cell.text.lower()
                for row in table.rows
                for cell in row.cells
            )

            if role_label.lower() not in table_text:
                continue

            for row in table.rows:
                if not any(
                    cell.text.strip().lower().startswith("signature:")
                    for cell in row.cells
                ):
                    continue

                for cell in row.cells:
                    drawings = cell._tc.xpath(
                        './/*[local-name()="drawing"]'
                    )

                    pictures = cell._tc.xpath(
                        './/*[local-name()="pict"]'
                    )

                    if drawings or pictures:
                        return True

                return False

        return False



    
    def get_location_checkboxes(
        self, possible_labels, table_index=None, row_index=None
    ):
        import re

        options = [
            "Telehealth",
            "Home",
            "School",
            "Clinic",
            "Community",
        ]

        normalized_labels = [
            " ".join(label.strip().lower().rstrip(":").split())
            for label in possible_labels
        ]

        for current_table_index, table in enumerate(self.document.tables):

            # If a specific table was requested, skip other tables.
            if (
                table_index is not None
                and current_table_index != table_index
            ):
                continue

            for current_row_index, row in enumerate(table.rows):

                # If a specific row was requested, skip other rows.
                if (
                    row_index is not None
                    and current_row_index != row_index
                ):
                    continue

                for cell in row.cells:
                    normalized_cell = " ".join(
                        cell.text.strip().lower().rstrip(":").split()
                    )

                    label_found = any(
                        normalized_cell.startswith(label)
                        for label in normalized_labels
                    )

                    if not label_found:
                        continue

                    # Location options may be in the label row
                    # or immediately following it.
                    for candidate_index in (
                        current_row_index,
                        current_row_index + 1,
                    ):
                        if candidate_index >= len(table.rows):
                            continue

                        for value_cell in table.rows[candidate_index].cells:
                            text = value_cell.text

                            # ----------------------------------
                            # METHOD 1: Unicode checkbox symbols
                            # ----------------------------------
                            selections = []

                            for option in options:
                                match = re.search(
                                    rf"([☒☑☐])\s*{re.escape(option)}\b",
                                    text,
                                    re.IGNORECASE,
                                )

                                if not match:
                                    selections = []
                                    break

                                selections.append(
                                    match.group(1) in ("☒", "☑")
                                )

                            if len(selections) == len(options):
                                return "  ".join(
                                    f"{'☒' if checked else '☐'} {option}"
                                    for option, checked in zip(
                                        options, selections
                                    )
                                )

                            # ----------------------------------
                            # METHOD 2: Word XML checkboxes
                            # ----------------------------------
                            checkboxes = value_cell._tc.xpath(
                                './/*[local-name()="checkbox"]'
                            )

                            if len(checkboxes) != len(options):
                                continue

                            states = []

                            for checkbox in checkboxes:
                                checked_elements = checkbox.xpath(
                                    './*[local-name()="checked"]'
                                )

                                checked_value = None

                                if checked_elements:
                                    checked_element = checked_elements[0]

                                    for name, value in checked_element.attrib.items():
                                        if (
                                            name.endswith("}val")
                                            or name == "val"
                                        ):
                                            checked_value = value
                                            break

                                states.append(
                                    checked_value in ("1", "true", "on")
                                )

                            return "  ".join(
                                f"{'☒' if checked else '☐'} {option}"
                                for option, checked in zip(
                                    options, states
                                )
                            )

        return None


    def debug_missing_spaces(self):

        print("DEBUG MISSING SPACES FUNCTION IS RUNNING")

        search_terms = [
            "theform",
            "the form",
            "mayfall",
            "may fall",
            "taken away",
            "takenaway",
        ]

        for table_index, table in enumerate(self.document.tables):
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):

                    text = cell.text

                    if not any(term in text.lower() for term in search_terms):
                        continue

                    print(f"\nTable: {table_index}")
                    print(f"Row: {row_index}")
                    print(f"Cell: {cell_index}")

                    print("\nCELL TEXT:")
                    print(repr(text))

                    print("\nPARAGRAPH RUNS:")

                    for paragraph in cell.paragraphs:
                        for run_index, run in enumerate(paragraph.runs):
                            print(
                                f"Run {run_index}: {repr(run.text)}"
                            )
